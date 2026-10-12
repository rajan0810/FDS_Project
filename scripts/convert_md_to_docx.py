import os
import re
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    # Add a single-cell table for clean background box effect
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    # Set thin border
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run = p.add_run(code_text.rstrip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    # Add spacing after table
    post_p = doc.add_paragraph()
    post_p.paragraph_format.space_before = Pt(0)
    post_p.paragraph_format.space_after = Pt(6)

def add_formatted_text(paragraph, text):
    """Parses inline bold, italic, code, and links."""
    # Pattern to match bold, italic, inline code, and links
    pattern = re.compile(
        r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    )
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif token.startswith('[') and ']' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                label, url = m.groups()
                run = paragraph.add_run(f"{label}")
                run.font.color.rgb = RGBColor(37, 99, 235)
                run.underline = True
            else:
                paragraph.add_run(token)
        else:
            paragraph.add_run(token)

def convert_md_to_docx(md_filepath: str, docx_filepath: str):
    doc = Document()
    
    # Set page margins (0.75 in)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
    with open(md_filepath, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    in_code_block = False
    code_block_lines = []
    code_lang = ""
    
    in_table = False
    table_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        
        # Check code blocks
        if stripped.startswith("```"):
            if not in_code_block:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_block_lines = []
            else:
                in_code_block = False
                add_code_block(doc, "".join(code_block_lines))
                code_block_lines = []
            i += 1
            continue
            
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
            
        # Check tables
        if stripped.startswith("|") and stripped.endswith("|"):
            # Table row
            cols = [c.strip() for c in stripped.split("|")[1:-1]]
            # Check if it's separator row (e.g. |:---|:---|)
            if all(re.match(r'^:?-+:?$', c) for c in cols):
                # separator row
                pass
            else:
                table_rows.append(cols)
            in_table = True
            i += 1
            continue
        elif in_table:
            # End of table, render it
            if table_rows:
                num_rows = len(table_rows)
                num_cols = max(len(r) for r in table_rows)
                t = doc.add_table(rows=num_rows, cols=num_cols)
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for r_idx, r_data in enumerate(table_rows):
                    for c_idx, val in enumerate(r_data):
                        cell = t.cell(r_idx, c_idx)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                        if r_idx == 0:
                            set_cell_background(cell, "F1F5F9")
                            run = p.add_run(val)
                            run.bold = True
                            run.font.size = Pt(9.5)
                        else:
                            add_formatted_text(p, val)
                            for r in p.runs:
                                r.font.size = Pt(9)
                # Border
                tblPr = t._tbl.tblPr
                tblBorders = parse_xml(f'''
                    <w:tblBorders {nsdecls("w")}>
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
                    </w:tblBorders>
                ''')
                tblPr.append(tblBorders)
                doc.add_paragraph()
            table_rows = []
            in_table = False
            
        # Check images
        img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', stripped)
        if img_match:
            caption, img_path_str = img_match.groups()
            img_path = Path(img_path_str)
            if not img_path.is_absolute():
                img_path = (Path(md_filepath).parent / img_path).resolve()
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(3)
                doc.add_picture(str(img_path), width=Inches(5.5))
                if caption:
                    cap_p = doc.add_paragraph()
                    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap_p.paragraph_format.space_before = Pt(2)
                    cap_p.paragraph_format.space_after = Pt(10)
                    cap_run = cap_p.add_run(f"Figure: {caption}")
                    cap_run.font.size = Pt(9)
                    cap_run.font.italic = True
                    cap_run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue
            
        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(6)
            add_formatted_text(h, stripped[2:])
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            add_formatted_text(h, stripped[3:])
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            add_formatted_text(h, stripped[4:])
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            add_formatted_text(h, stripped[5:])
        elif stripped.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, stripped[2:])
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(30, 64, 175)
        elif stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_formatted_text(p, stripped[2:])
        elif re.match(r'^\d+\.\s', stripped):
            m = re.match(r'^\d+\.\s(.*)$', stripped)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            add_formatted_text(p, m.group(1))
        elif stripped.startswith("---") or stripped.startswith("==="):
            # Horizontal rule
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run("―" * 45)
            r.font.color.rgb = RGBColor(203, 213, 225)
        elif stripped:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            add_formatted_text(p, stripped)
            
        i += 1
        
    doc.save(docx_filepath)
    print(f"Successfully generated DOCX: {docx_filepath}")

if __name__ == "__main__":
    md_file = sys.argv[1] if len(sys.argv) > 1 else "project_submission_report.md"
    docx_file = sys.argv[2] if len(sys.argv) > 2 else "project_submission_report.docx"
    convert_md_to_docx(md_file, docx_file)
