import os
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run = p.add_run(code_text.rstrip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    post_p = doc.add_paragraph()
    post_p.paragraph_format.space_before = Pt(0)
    post_p.paragraph_format.space_after = Pt(8)

def main():
    doc = Document()
    
    # Page setup (0.75 inch margins)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(12)
    title_run = title_p.add_run("Python Source Code Documentation")
    title_run.bold = True
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.color.rgb = RGBColor(17, 24, 39)
    
    # List of files and their detailed explanations
    files_data = [
        {
            "filename": "main.py",
            "path": "main.py",
            "explanation": "This is the primary command-line interface (CLI) and entry point of the project. It provides subcommands to run various pipeline stages: `setup` (validates the environment, API key, and datasets), `inspect` (examines raw QASPER samples, reference answers, and evidence), `eda` (executes the full exploratory data analysis suite: context length, relevance, evidence position, redundancy, and cross-document similarity), `demo` (runs an end-to-end adaptive context selection pipeline including reranking, comparing baseline vs. compressed LLM generation), `interactive` (allows users to ask custom questions on loaded papers), and `experiment` (runs the Context-vs-Accuracy / Latency / Cost-Quality Tradeoff sweep across multiple context-size budgets and sampled questions). It also seeds Python's and NumPy's random number generators from `config.yaml`'s `RANDOM_SEED` at startup for reproducibility."
        },
        {
            "filename": "src/config.py",
            "path": "src/config.py",
            "explanation": "This module handles environment configuration and parameter loading. It reads runtime hyperparameters (such as chunk token sizes, embedding models, and redundancy thresholds) from `config/config.yaml` and manages secure access to the Gemini API key and model name from the `.env` file."
        },
        {
            "filename": "src/dataset.py",
            "path": "src/dataset.py",
            "explanation": "This module is responsible for loading and caching the scientific QASPER dataset. It automatically downloads the official dataset tarball from AWS S3 if not present, extracts the validation split JSON, and caches it locally under `data/cache/` to ensure fast, offline subsequent runs."
        },
        {
            "filename": "src/preprocessing.py",
            "path": "src/preprocessing.py",
            "explanation": "This module implements data structure normalization and preprocessing routines. It defines the `PaperRecord` and `QuestionRecord` classes, flattens hierarchical research papers into structured section/paragraph dictionaries with unique paragraph IDs and token counts, and extracts questions along with their reference answers and gold-annotated evidence."
        },
        {
            "filename": "src/chunking.py",
            "path": "src/chunking.py",
            "explanation": "This module implements token-aware document chunking. It aggregates consecutive paragraphs into fixed-size chunks (e.g., 500 tokens) with a sliding overlap window (e.g., 50 tokens) to preserve cross-paragraph context boundaries and stores exact paragraph mappings for each chunk."
        },
        {
            "filename": "src/embeddings.py",
            "path": "src/embeddings.py",
            "explanation": "This module manages dense vector embeddings using Hugging Face's `sentence-transformers` (specifically `all-MiniLM-L6-v2`). It converts textual chunks into 384-dimensional dense vectors and provides utilities to save and load precomputed embeddings to/from the `data/embeddings/` directory."
        },
        {
            "filename": "src/retrieval.py",
            "path": "src/retrieval.py",
            "explanation": "This module implements vector search and retrieval functionality (the Candidate Retrieval pipeline stage). It calculates cosine similarity between the question vector and all document chunk vectors, ranks candidate chunks by similarity score, and maps ground-truth gold evidence strings back to chunk IDs to enable retrieval precision/recall evaluation."
        },
        {
            "filename": "src/reranking.py",
            "path": "src/reranking.py",
            "explanation": "This module implements the Context Reranking pipeline stage. It wraps a `sentence-transformers` cross-encoder (`cross-encoder/ms-marco-MiniLM-L-6-v2`) that jointly scores each (query, chunk) pair, which is more accurate than the bi-encoder similarity score used for initial retrieval, and re-sorts the retrieved candidates before adaptive selection."
        },
        {
            "filename": "src/redundancy.py",
            "path": "src/redundancy.py",
            "explanation": "This module implements the Redundancy Removal pipeline stage as well as the pairwise redundancy analysis used in EDA. `flag_redundant_chunks` greedily walks relevance-ranked candidates and drops any chunk whose cosine similarity to an already-kept chunk exceeds the configured redundancy threshold (0.85). `calculate_pairwise_redundancy` computes the full pairwise similarity matrix for analysis/visualization purposes."
        },
        {
            "filename": "src/adaptive_selection.py",
            "path": "src/adaptive_selection.py",
            "explanation": "This module implements the Adaptive Context Selection and Context Ordering pipeline stages as three explicit, composable functions: `select_by_relevance` (greedy relevance-based cutoff with min/max chunk budgets), `remove_redundant_chunks` (delegates to `src/redundancy.py`), and `order_context` (arranges the final evidence chunks in original document order to mitigate 'lost in the middle' effects). `adaptive_select` chains all three for convenience."
        },
        {
            "filename": "src/generation.py",
            "path": "src/generation.py",
            "explanation": "This module interfaces with Google Gemini (`gemini-3.1-pro-preview`) via the `google-genai` SDK (the LLM Generation pipeline stage). It formats grounded question-answering prompts strictly constraining the model to use only the provided context, invokes the API, tracks inference latency, and extracts real prompt/completion/total token usage from the API response for cost accounting."
        },
        {
            "filename": "src/evaluation.py",
            "path": "src/evaluation.py",
            "explanation": "This module provides the full evaluation-metric suite. Answer metrics: token-level F1, Exact Match (EM), ROUGE-1/2/L (via the `rouge-score` package), and a Groundedness score (a hybrid of answer/context token overlap and answer/context embedding cosine similarity). Retrieval metrics: Precision@K, Recall@K, MRR, and NDCG@K (binary-relevance, log2-discounted). System/cost metrics: reusable `compression_ratio` and `throughput` helpers."
        },
        {
            "filename": "src/eda.py",
            "path": "src/eda.py",
            "explanation": "This module contains the full Exploratory Data Analysis (EDA) suite: Context Length Analysis (document/chunk token distributions), Relevance Analysis (per-query retrieval usefulness), Evidence Position Analysis ('lost-in-the-middle' study), Redundancy Analysis (corpus-wide redundancy rate plus a representative similarity heatmap), and Cross-Document Similarity Analysis (pools each paper's chunk embeddings and computes a paper-vs-paper similarity matrix/heatmap across the sampled corpus)."
        },
        {
            "filename": "src/experiments.py",
            "path": "src/experiments.py",
            "explanation": "This module runs the LLM-driven experiment sweep behind the Context-vs-Accuracy, Latency, and Cost-Quality Tradeoff analyses. For a sample of QASPER questions, it runs the full retrieve -> rerank -> select -> deduplicate -> order -> generate pipeline at several context-size budgets (plus a full top-K baseline with no adaptive filtering), recording answer quality (F1/EM/ROUGE/Groundedness), token usage, latency, and memory delta per run, then aggregates a per-configuration cost-quality tradeoff table and plots."
        },
        {
            "filename": "src/utils.py",
            "path": "src/utils.py",
            "explanation": "This module provides shared helper functions, including a fallback token estimation counter based on word count heuristics."
        }
    ]
    
    for item in files_data:
        file_path = Path(item["path"])
        if not file_path.exists():
            print(f"Warning: {file_path} not found")
            continue
            
        with open(file_path, "r", encoding="utf-8") as f:
            code_content = f.read()
            
        # File Heading
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        h_run = h.add_run(f"File: {item['filename']}")
        h_run.bold = True
        h_run.font.name = "Calibri"
        h_run.font.size = Pt(14)
        h_run.font.color.rgb = RGBColor(30, 58, 138)
        
        # Explanation
        exp_p = doc.add_paragraph()
        exp_p.paragraph_format.space_before = Pt(2)
        exp_p.paragraph_format.space_after = Pt(6)
        exp_p.paragraph_format.line_spacing = 1.15
        
        lbl_run = exp_p.add_run("Explanation: ")
        lbl_run.bold = True
        lbl_run.font.name = "Calibri"
        lbl_run.font.size = Pt(11)
        
        exp_run = exp_p.add_run(item["explanation"])
        exp_run.font.name = "Calibri"
        exp_run.font.size = Pt(11)
        exp_run.font.color.rgb = RGBColor(55, 65, 81)
        
        # Code block
        add_code_block(doc, code_content)
        
    out_docx = "Source_Code_Submission.docx"
    doc.save(out_docx)
    print(f"Successfully generated {out_docx}")

if __name__ == "__main__":
    main()
